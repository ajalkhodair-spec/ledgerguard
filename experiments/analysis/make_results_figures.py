#!/usr/bin/env python3
"""Generate the publication Results figures from the frozen LedgerGuard workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import random
import shutil
import subprocess
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

os.environ.setdefault("MPLCONFIGDIR", "/tmp/ledgerguard-results-figures-matplotlib")

import matplotlib as mpl
mpl.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from openpyxl import load_workbook
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[2]
SEED = 20260828
RESAMPLES = 10_000
TOLERANCE = 1e-9
SCRIPT_PATH = Path(__file__).resolve()

BLUE = "#0072B2"
ORANGE = "#D55E00"
GREEN = "#009E73"
PURPLE = "#7B3294"
SKY = "#56B4E9"
GRAY = "#5B6573"
LIGHT_GRAY = "#D7DCE2"
BLACK = "#222222"
WHITE = "#FFFFFF"

MODE_LABELS = {
    "naive_per_device": "Per-device reporting",
    "v1_integrity_root": "Integrity-only aggregation",
    "v2_complete_witnessed": "Complete-cohort witnessed aggregation",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def percentile(values: list[float], probability: float) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def median(values: list[float]) -> float:
    return float(np.median(np.asarray(values, dtype=float)))


def mean(values: list[float]) -> float:
    return float(np.mean(np.asarray(values, dtype=float)))


def assert_close(actual: float, expected: float, label: str, tolerance: float = TOLERANCE) -> None:
    if not math.isclose(actual, expected, rel_tol=tolerance, abs_tol=tolerance):
        raise ValueError(f"{label}: recomputed={actual!r}, workbook={expected!r}")


def select_font() -> str:
    candidates = ["Times New Roman", "Tinos", "Liberation Serif"]
    installed = {font.name for font in font_manager.fontManager.ttflist}
    for candidate in candidates:
        if candidate in installed:
            return candidate
    raise RuntimeError("No approved Times-compatible font is installed")


def configure_style(font: str) -> None:
    mpl.rcParams.update({
        "font.family": font,
        "font.size": 8.5,
        "axes.titlesize": 9,
        "axes.labelsize": 9,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "legend.fontsize": 8,
        "axes.linewidth": 0.9,
        "lines.linewidth": 1.6,
        "lines.markersize": 6,
        "figure.facecolor": WHITE,
        "axes.facecolor": WHITE,
        "savefig.facecolor": WHITE,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "svg.fonttype": "none",
        "axes.unicode_minus": False,
    })


def style_axis(ax: plt.Axes, grid_axis: str | None = "y") -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    for spine in ax.spines.values():
        spine.set_linewidth(0.9)
    ax.tick_params(width=0.8, length=3.5)
    if grid_axis:
        ax.grid(axis=grid_axis, color=LIGHT_GRAY, linewidth=0.6, alpha=0.7)
        ax.set_axisbelow(True)


def panel_label(ax: plt.Axes, label: str) -> None:
    ax.text(-0.13, 1.06, label, transform=ax.transAxes, fontweight="bold", fontsize=9,
            ha="left", va="bottom", clip_on=False)


class Sources:
    def __init__(self, workbook_path: Path):
        self.workbook_path = workbook_path.resolve()
        self.workbook_hash = sha256(self.workbook_path)
        self.workbook = load_workbook(self.workbook_path, read_only=False, data_only=True)
        self.index = {row["worksheet"]: row for row in self.sheet("Result Index")}
        manifest_path = ROOT / "release/evidence_manifest.json"
        self.manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.manifest_files = {item["path"]: item for item in self.manifest["files"]}

    def sheet(self, name: str) -> list[dict[str, Any]]:
        if name not in self.workbook.sheetnames:
            raise KeyError(f"Missing canonical workbook sheet: {name}")
        worksheet = self.workbook[name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value) for value in rows[0] if value is not None]
        return [dict(zip(headers, row[:len(headers)], strict=True)) for row in rows[1:] if any(value is not None for value in row)]

    def csv_path(self, sheet_name: str) -> Path:
        indexed = self.index.get(sheet_name)
        if not indexed:
            raise KeyError(f"No Result Index entry for {sheet_name}")
        source_path = str(indexed["source_path"])
        path = ROOT / source_path
        if not path.is_file():
            raise FileNotFoundError(path)
        return path

    def csv_rows(self, sheet_name: str) -> tuple[Path, list[dict[str, str]]]:
        path = self.csv_path(sheet_name)
        with path.open(newline="", encoding="utf-8") as stream:
            rows = list(csv.DictReader(stream))
        expected = int(self.index[sheet_name]["data_rows"])
        if len(rows) != expected:
            raise ValueError(f"{sheet_name}: CSV rows={len(rows)}, Result Index rows={expected}")
        return path, rows

    def file_hash(self, relative_path: str) -> str | None:
        item = self.manifest_files.get(relative_path)
        return item.get("sha256") if item else None


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"No plot data for {path.name}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def save_figure(fig: plt.Figure, base: Path) -> list[Path]:
    base.parent.mkdir(parents=True, exist_ok=True)
    outputs = [base.with_suffix(".pdf"), base.with_suffix(".svg"), base.with_suffix(".png")]
    fig.savefig(outputs[0], bbox_inches="tight", pad_inches=0.04, metadata={"Creator": "LedgerGuard results figure pipeline"})
    fig.savefig(outputs[1], bbox_inches="tight", pad_inches=0.04, metadata={"Creator": "LedgerGuard results figure pipeline"})
    fig.savefig(outputs[2], dpi=600, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)
    return outputs


def output_records(paths: list[Path]) -> list[dict[str, str]]:
    return [{"path": str(path.relative_to(ROOT)), "sha256": sha256(path)} for path in paths]


def metadata_base(name: str, label: str, evidence_class: str, sources: Sources,
                  sheets: list[str], source_paths: list[Path], columns: list[str],
                  filters: dict[str, Any], aggregation: str, row_count: int,
                  excluded: int, plotted: list[float], outputs: list[Path], font: str) -> dict[str, Any]:
    return {
        "figure_semantic_name": name,
        "manuscript_label": label,
        "evidence_class": evidence_class,
        "workbook_filename": sources.workbook_path.name,
        "workbook_sha256": sources.workbook_hash,
        "source_worksheet_names": sheets,
        "source_csv_paths": [str(path.relative_to(ROOT)) for path in source_paths],
        "source_csv_sha256": {
            str(path.relative_to(ROOT)): sources.file_hash(str(path.relative_to(ROOT)))
            for path in source_paths if sources.file_hash(str(path.relative_to(ROOT))) is not None
        },
        "source_columns": columns,
        "filters": filters,
        "aggregation_method": aggregation,
        "bootstrap_seed": SEED if "bootstrap" in aggregation.lower() else None,
        "row_count": row_count,
        "excluded_row_count": excluded,
        "minimum": min(plotted) if plotted else None,
        "median": median(plotted) if plotted else None,
        "maximum": max(plotted) if plotted else None,
        "output_files": output_records(outputs),
        "plotting_script_path": str(SCRIPT_PATH.relative_to(ROOT)),
        "plotting_script_sha256": sha256(SCRIPT_PATH),
        "selected_font": font,
        "generated_timestamp": datetime.now(timezone.utc).isoformat(),
        "git_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "validation_status": "PENDING",
    }


def render_block_period(ax: plt.Axes, grouped: dict[int, list[float]], annotate: bool = True) -> None:
    periods = [1, 2, 4]
    values = [grouped[period] for period in periods]
    box = ax.boxplot(values, positions=range(3), widths=0.48, patch_artist=True, showfliers=False,
                     medianprops={"color": BLACK, "linewidth": 1.5},
                     whiskerprops={"color": GRAY}, capprops={"color": GRAY})
    for patch in box["boxes"]:
        patch.set_facecolor(SKY)
        patch.set_alpha(0.28)
        patch.set_edgecolor(BLUE)
    rng = np.random.default_rng(SEED)
    for index, group in enumerate(values):
        jitter = rng.uniform(-0.13, 0.13, len(group))
        ax.scatter(index + jitter, group, s=18, marker="o", facecolor=WHITE, edgecolor=BLUE,
                   linewidth=0.75, alpha=0.9, zorder=3)
        med = median(group)
        ax.text(index, med * 0.86, f"Median {med:.3f} s", ha="center", va="top", fontsize=7.5)
    if annotate:
        outlier = max(grouped[1])
        ax.annotate(f"Retained observation\n{outlier:.6f} s", xy=(0, outlier), xytext=(0.38, outlier * 0.72),
                    arrowprops={"arrowstyle": "-", "color": GRAY, "linewidth": 0.8}, fontsize=7.5,
                    ha="left", va="top")
    ax.set_xticks(range(3), ["1", "2", "4"])
    ax.set_xlabel("Configured block period")
    ax.set_ylabel("Submission-to-receipt latency (s, log scale)")
    ax.set_yscale("log")
    ax.set_ylim(0.75, 45)
    style_axis(ax, "y")


def figure_r1(sources: Sources, out: Path, font: str) -> dict[str, Any]:
    source_path, rows = sources.csv_rows("Block Period Raw")
    summary = {int(row["block_period_seconds"]): row for row in sources.sheet("Block Period Summary")}
    valid = [row for row in rows if row["status"] == "success"]
    grouped = {period: [float(row["latency_seconds"]) for row in valid if int(row["block_period_seconds"]) == period] for period in [1, 2, 4]}
    if [len(grouped[p]) for p in [1, 2, 4]] != [30, 30, 30]:
        raise ValueError("Block-period profiles do not contain 30 valid observations each")
    for period, values in grouped.items():
        assert_close(median(values), float(summary[period]["median_seconds"]), f"block-period {period}s median")
        assert_close(min(values), float(summary[period]["minimum_seconds"]), f"block-period {period}s minimum")
        assert_close(max(values), float(summary[period]["maximum_seconds"]), f"block-period {period}s maximum")
    fig, ax = plt.subplots(figsize=(5.4, 3.6))
    render_block_period(ax, grouped)
    outputs = save_figure(fig, out / "fig_results_block_period_sensitivity")
    plot_rows = [{"configured_block_period_seconds": int(row["block_period_seconds"]), "repetition": int(row["repetition"]),
                  "operation": row["operation"], "submission_to_receipt_latency_seconds": float(row["latency_seconds"]),
                  "receipt_poll_interval_seconds": float(row["receipt_poll_interval_seconds"]), "status": row["status"]} for row in valid]
    write_csv(out / "plot_data/fig_results_block_period_sensitivity.csv", plot_rows)
    return metadata_base("fig_results_block_period_sensitivity", "fig:block-period-sensitivity",
                         "representative local Besu state-changing transaction sensitivity", sources,
                         ["Block Period Raw", "Block Period Summary"], [source_path],
                         list(plot_rows[0]), {"status": "success", "block_period_seconds": [1, 2, 4]},
                         "All observations with box summaries; no P95 reported for n=30", len(valid), len(rows) - len(valid),
                         [value for group in grouped.values() for value in group], outputs, font)


def fleet_bootstrap_summaries(rows: list[dict[str, str]]) -> dict[tuple[int, float], dict[str, float]]:
    grouped: dict[tuple[int, float], list[float]] = defaultdict(list)
    for row in rows:
        grouped[(int(row["fleet_size"]), float(row["configured_failure_rate"]))].append(float(row["final_adoption_rate"]))
    rng = random.Random(SEED)
    results = {}
    for key, values in sorted(grouped.items()):
        boot = [mean(rng.choices(values, k=len(values))) for _ in range(RESAMPLES)]
        results[key] = {"mean": mean(values), "low": percentile(boot, 0.025), "high": percentile(boot, 0.975)}
    return results


def render_fleet_success(ax: plt.Axes, points: dict[tuple[int, float], dict[str, float]]) -> None:
    sizes = [100, 500, 1000]
    rates = [0.01, 0.02, 0.05]
    colors = [BLUE, ORANGE, GREEN]
    markers = ["o", "s", "^"]
    offsets = [-0.18, 0, 0.18]
    for rate, color, marker, offset in zip(rates, colors, markers, offsets, strict=True):
        center = np.arange(3, dtype=float) + offset
        values = np.array([100 * points[(size, rate)]["mean"] for size in sizes])
        low = np.array([100 * points[(size, rate)]["low"] for size in sizes])
        high = np.array([100 * points[(size, rate)]["high"] for size in sizes])
        ax.errorbar(center, values, yerr=np.vstack((values - low, high - values)), fmt=marker,
                    color=color, markerfacecolor=WHITE, markeredgewidth=1.2, linewidth=0,
                    elinewidth=1.4, capsize=3, label=f"Configured failure probability: {rate * 100:.0f}%")
    ax.set_xticks(range(3), ["100", "500", "1,000"])
    ax.set_xlim(-0.5, 2.5)
    ax.set_ylim(93, 100.2)
    ax.set_xlabel("Cohort size (software identities)")
    ax.set_ylabel("Mean success share (%)")
    ax.legend(frameon=False, ncol=1, loc="lower left")
    style_axis(ax, "y")


def figure_r2(sources: Sources, out: Path, font: str) -> dict[str, Any]:
    source_path, rows = sources.csv_rows("Fleet Runs")
    workbook_summary = {(int(row["fleet_size"]), float(row["configured_failure_rate"])): row for row in sources.sheet("Fleet Summary")}
    points = fleet_bootstrap_summaries(rows)
    if len(points) != 9 or len(rows) != 180:
        raise ValueError("Fleet campaign must contain nine configurations and 180 runs")
    for key, result in points.items():
        expected = workbook_summary[key]
        assert_close(result["mean"], float(expected["mean_adoption_rate"]), f"fleet {key} mean")
        assert_close(result["low"], float(expected["mean_ci95_low"]), f"fleet {key} CI low")
        assert_close(result["high"], float(expected["mean_ci95_high"]), f"fleet {key} CI high")
    fig, ax = plt.subplots(figsize=(5.4, 3.5))
    render_fleet_success(ax, points)
    outputs = save_figure(fig, out / "fig_results_success_share_software_cohorts")
    plot_rows = [{"cohort_size_software_identities": size, "configured_failure_probability": rate,
                  "seed_count": 20, "mean_success_share_percent": 100 * points[(size, rate)]["mean"],
                  "bootstrap_95_ci_low_percent": 100 * points[(size, rate)]["low"],
                  "bootstrap_95_ci_high_percent": 100 * points[(size, rate)]["high"]}
                 for size in [100, 500, 1000] for rate in [0.01, 0.02, 0.05]]
    write_csv(out / "plot_data/fig_results_success_share_software_cohorts.csv", plot_rows)
    return metadata_base("fig_results_success_share_software_cohorts", "fig:fleet-success-share",
                         "authenticated software-device execution", sources, ["Fleet Runs", "Fleet Summary"],
                         [source_path], list(plot_rows[0]), {}, "Mean and 10,000-resample bootstrap 95% confidence interval",
                         len(rows), 0, [row["mean_success_share_percent"] for row in plot_rows], outputs, font)


OUTCOME_FIELDS = {"FAIL": "fail_count", "MISSING": "missing_count", "REJECTED": "rejected_count",
                  "ROLLBACK": "rollback_count", "SUCCESS": "success_count"}


def outcome_bootstrap_summaries(rows: list[dict[str, str]]) -> dict[tuple[int, float, str], dict[str, float]]:
    grouped: dict[tuple[int, float, str], list[float]] = defaultdict(list)
    for row in rows:
        expected = int(row["expected_count"])
        counts = {outcome: int(row[field]) for outcome, field in OUTCOME_FIELDS.items()}
        if sum(counts.values()) != expected:
            raise ValueError(f"Outcome counts do not reconcile for {row['run_id']}")
        for outcome, count in counts.items():
            grouped[(int(row["fleet_size"]), float(row["configured_failure_rate"]), outcome)].append(100 * count / expected)
    rng = random.Random(SEED)
    result = {}
    for key, values in sorted(grouped.items()):
        boot = [mean(rng.choices(values, k=len(values))) for _ in range(RESAMPLES)]
        result[key] = {"mean": mean(values), "low": percentile(boot, 0.025), "high": percentile(boot, 0.975)}
    return result


def render_outcome(ax: plt.Axes, outcome: str, values: list[dict[str, float]], panel: str) -> None:
    rates = np.array([1, 2, 5], dtype=float)
    mean_values = np.array([item["mean"] for item in values])
    low = np.array([item["low"] for item in values])
    high = np.array([item["high"] for item in values])
    color = {"SUCCESS": GREEN, "ROLLBACK": ORANGE, "FAIL": PURPLE}[outcome]
    marker = {"SUCCESS": "o", "ROLLBACK": "s", "FAIL": "^"}[outcome]
    ax.errorbar(rates, mean_values, yerr=np.vstack((mean_values - low, high - mean_values)),
                color=color, marker=marker, markerfacecolor=WHITE, markeredgewidth=1.1,
                linestyle={"SUCCESS": "-", "ROLLBACK": "--", "FAIL": "-."}[outcome], capsize=3)
    ax.set_xticks(rates)
    ax.set_xlim(0.6, 5.4)
    ax.set_ylim({"SUCCESS": (93, 100), "ROLLBACK": (0, 2.1), "FAIL": (0, 4.5)}[outcome])
    ax.set_xlabel("Configured failure probability (%)")
    ax.set_ylabel("Mean outcome share (%)")
    ax.set_title(outcome.title())
    panel_label(ax, panel)
    style_axis(ax, "y")


def figure_r3(sources: Sources, out: Path, font: str) -> dict[str, Any]:
    source_path, rows = sources.csv_rows("Fleet Runs")
    workbook_outcomes = {(int(row["fleet_size"]), float(row["configured_failure_rate"]), str(row["outcome"])): row for row in sources.sheet("Fleet Outcomes")}
    all_results = outcome_bootstrap_summaries(rows)
    outcomes = ["SUCCESS", "ROLLBACK", "FAIL"]
    selected = {(rate, outcome): all_results[(500, rate, outcome)] for rate in [0.01, 0.02, 0.05] for outcome in outcomes}
    for (rate, outcome), result in selected.items():
        expected = workbook_outcomes[(500, rate, outcome)]
        assert_close(result["mean"], float(expected["mean"]), f"{outcome} {rate} mean")
        assert_close(result["low"], float(expected["mean_ci95_low"]), f"{outcome} {rate} CI low")
        assert_close(result["high"], float(expected["mean_ci95_high"]), f"{outcome} {rate} CI high")
    outputs: list[Path] = []
    for index, outcome in enumerate(outcomes):
        fig, ax = plt.subplots(figsize=(3.45, 2.9))
        render_outcome(ax, outcome, [selected[(rate, outcome)] for rate in [0.01, 0.02, 0.05]], f"({chr(97 + index)})")
        outputs += save_figure(fig, out / f"panels/fig_results_failure_probability_sensitivity_{chr(97 + index)}")
    fig, axes = plt.subplots(1, 3, figsize=(7.2, 2.75), layout="constrained")
    for index, (ax, outcome) in enumerate(zip(axes, outcomes, strict=True)):
        render_outcome(ax, outcome, [selected[(rate, outcome)] for rate in [0.01, 0.02, 0.05]], f"({chr(97 + index)})")
    outputs += save_figure(fig, out / "fig_results_failure_probability_sensitivity")
    plot_rows = [{"cohort_size_software_identities": 500, "configured_failure_probability": rate,
                  "outcome": outcome, "seed_count": 20, "mean_outcome_share_percent": selected[(rate, outcome)]["mean"],
                  "bootstrap_95_ci_low_percent": selected[(rate, outcome)]["low"],
                  "bootstrap_95_ci_high_percent": selected[(rate, outcome)]["high"]}
                 for outcome in outcomes for rate in [0.01, 0.02, 0.05]]
    write_csv(out / "plot_data/fig_results_failure_probability_sensitivity.csv", plot_rows)
    relevant_runs = [row for row in rows if int(row["fleet_size"]) == 500]
    return metadata_base("fig_results_failure_probability_sensitivity", "fig:failure-probability-sensitivity",
                         "authenticated software-device execution", sources, ["Fleet Outcomes", "Fleet Runs"],
                         [source_path], list(plot_rows[0]), {"cohort_size": 500, "outcomes": outcomes},
                         "Outcome count divided by expected cohort; mean and 10,000-resample bootstrap 95% confidence interval",
                         len(relevant_runs), 0, [row["mean_outcome_share_percent"] for row in plot_rows], outputs, font)


def render_network(ax: plt.Axes, grouped: dict[str, list[float]], profile_labels: dict[str, str]) -> None:
    profiles = ["P1", "P2", "P3"]
    values = [grouped[profile] for profile in profiles]
    positions = np.arange(3)
    box = ax.boxplot(values, positions=positions, orientation="horizontal", widths=0.46, patch_artist=True,
                     showfliers=False, medianprops={"color": BLACK, "linewidth": 1.5})
    for patch, color in zip(box["boxes"], [SKY, BLUE, PURPLE], strict=True):
        patch.set_facecolor(color)
        patch.set_alpha(0.25)
        patch.set_edgecolor(color)
    rng = np.random.default_rng(SEED)
    for index, (group, color, marker) in enumerate(zip(values, [SKY, BLUE, PURPLE], ["o", "s", "^"], strict=True)):
        jitter = rng.uniform(-0.13, 0.13, len(group))
        ax.scatter(group, index + jitter, s=18, marker=marker, facecolor=WHITE, edgecolor=color,
                   linewidth=0.8, alpha=0.95, zorder=3)
    maximum = max(grouped["P3"])
    ax.annotate(f"Retained observation: {maximum:.6f} s", xy=(maximum, 2), xytext=(maximum / 14, 1.55),
                arrowprops={"arrowstyle": "-", "color": GRAY, "linewidth": 0.8}, fontsize=7.5,
                ha="right", va="center")
    ax.set_yticks(positions, [profile_labels[profile] for profile in profiles])
    ax.invert_yaxis()
    ax.set_xscale("log")
    ax.set_xlim(0.65, 1450)
    ax.set_xlabel("End-to-end artifact retrieval time (s, log scale)")
    style_axis(ax, "x")


def figure_r4(sources: Sources, out: Path, font: str) -> dict[str, Any]:
    summary_rows = sources.sheet("Network Summary")
    distribution = {str(row["profile"]): row for row in sources.sheet("Network Distribution")}
    raw_relative = str(summary_rows[0]["raw_evidence_path"])
    raw_path = ROOT / raw_relative
    rows = [json.loads(line) for line in raw_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    grouped = {profile: [float(row["retrieval_time_seconds"]) for row in rows if row["profile"] == profile] for profile in ["P1", "P2", "P3"]}
    if [len(grouped[p]) for p in ["P1", "P2", "P3"]] != [20, 20, 20]:
        raise ValueError("Network profiles do not contain 20 observations each")
    for profile, values in grouped.items():
        expected = distribution[profile]
        assert_close(min(values), float(expected["minimum"]), f"network {profile} minimum")
        assert_close(median(values), float(expected["median"]), f"network {profile} median")
        assert_close(max(values), float(expected["maximum"]), f"network {profile} maximum")
    labels = {
        "P1": "P1 | 20 ms application delay\n5 Mbps read limit",
        "P2": "P2 | 80 ms application delay\n5 Mbps read limit",
        "P3": "P3 | 200 ms application delay\n0.256 Mbps read limit",
    }
    fig, ax = plt.subplots(figsize=(7.2, 3.35))
    render_network(ax, grouped, labels)
    outputs = save_figure(fig, out / "fig_results_application_shaped_ipfs_retrieval")
    plot_rows = [{"profile": row["profile"], "repetition": int(row["repetition"]),
                  "configured_application_delay_ms": float(row["configured_application_delay_ms"]),
                  "configured_read_rate_limit_mbps": float(row["configured_read_rate_limit_mbps"]),
                  "artifact_size_bytes": int(row["artifact_size_bytes"]),
                  "retrieval_time_seconds": float(row["retrieval_time_seconds"]), "status": row["status"]} for row in rows]
    write_csv(out / "plot_data/fig_results_application_shaped_ipfs_retrieval.csv", plot_rows)
    return metadata_base("fig_results_application_shaped_ipfs_retrieval", "fig:ipfs-retrieval",
                         "application-shaped local IPFS retrieval", sources, ["Network Summary", "Network Distribution"],
                         [raw_path], list(plot_rows[0]), {}, "All observations with box summaries on a logarithmic axis",
                         len(rows), 0, [row["retrieval_time_seconds"] for row in plot_rows], outputs, font)


def render_cache_origin(ax: plt.Axes, medians: dict[str, float], panel: str) -> None:
    modes = ["OFF", "ON"]
    values = [medians[mode] / (1024 * 1024) for mode in modes]
    y = np.arange(2)
    ax.barh(y, values, height=0.48, color=[ORANGE, GREEN], edgecolor=BLACK, linewidth=0.7)
    ax.set_yticks(y, ["Cache OFF", "Cache ON"])
    ax.invert_yaxis()
    ax.set_xscale("log")
    ax.set_xlim(0.35, 900)
    ax.set_xlabel("Median origin-side transfer (MiB)")
    ax.set_title("Origin-side transfer")
    panel_label(ax, panel)
    style_axis(ax, "x")


def render_cache_time(ax: plt.Axes, grouped: dict[str, list[float]], panel: str) -> None:
    modes = ["OFF", "ON"]
    values = [grouped[mode] for mode in modes]
    box = ax.boxplot(values, positions=[0, 1], widths=0.48, patch_artist=True, showfliers=False,
                     medianprops={"color": BLACK, "linewidth": 1.5})
    for patch, color in zip(box["boxes"], [ORANGE, GREEN], strict=True):
        patch.set_facecolor(color)
        patch.set_alpha(0.25)
        patch.set_edgecolor(color)
    rng = np.random.default_rng(SEED + 5)
    for index, (group, color, marker) in enumerate(zip(values, [ORANGE, GREEN], ["s", "o"], strict=True)):
        ax.scatter(index + rng.uniform(-0.12, 0.12, len(group)), group, s=18, marker=marker,
                   facecolor=WHITE, edgecolor=color, linewidth=0.8, zorder=3)
    ax.set_xticks([0, 1], ["Cache OFF", "Cache ON"])
    ax.set_ylabel("Total retrieval time per trial (s)")
    ax.set_title("Total retrieval time")
    panel_label(ax, panel)
    style_axis(ax, "y")


def figure_r5(sources: Sources, out: Path, font: str) -> dict[str, Any]:
    source_path, rows = sources.csv_rows("Cache Trials")
    ablation = {str(row["cache_mode"]): row for row in sources.sheet("Cache Ablation")}
    grouped = {mode: [float(row["total_retrieval_time_seconds"]) for row in rows if row["cache_mode"] == mode] for mode in ["OFF", "ON"]}
    medians = {mode: median([float(row["origin_bytes_transferred"]) for row in rows if row["cache_mode"] == mode]) for mode in ["OFF", "ON"]}
    if [len(grouped[mode]) for mode in ["OFF", "ON"]] != [10, 10]:
        raise ValueError("Cache modes do not contain 10 independent trials each")
    for mode in ["OFF", "ON"]:
        assert_close(medians[mode], float(ablation[mode]["median_origin_bytes_transferred"]), f"cache {mode} origin bytes")
        assert_close(median(grouped[mode]), float(ablation[mode]["median_total_retrieval_time_seconds"]), f"cache {mode} time")
    outputs: list[Path] = []
    fig, ax = plt.subplots(figsize=(3.45, 2.75)); render_cache_origin(ax, medians, "(a)")
    outputs += save_figure(fig, out / "panels/fig_results_local_cache_reuse_a")
    fig, ax = plt.subplots(figsize=(3.45, 2.75)); render_cache_time(ax, grouped, "(b)")
    outputs += save_figure(fig, out / "panels/fig_results_local_cache_reuse_b")
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.0), layout="constrained")
    render_cache_origin(axes[0], medians, "(a)"); render_cache_time(axes[1], grouped, "(b)")
    outputs += save_figure(fig, out / "fig_results_local_cache_reuse")
    plot_rows = []
    for row in rows:
        plot_rows.append({"record_type": "trial", "cache_mode": row["cache_mode"], "trial": int(row["trial"]),
                          "artifact_requests": int(row["artifact_requests"]), "origin_requests": int(row["origin_requests"]),
                          "cache_hits": int(row["cache_hits"]), "origin_bytes_transferred": int(row["origin_bytes_transferred"]),
                          "origin_side_transfer_mib": float(row["origin_bytes_transferred"]) / (1024 * 1024),
                          "total_retrieval_time_seconds": float(row["total_retrieval_time_seconds"]),
                          "integrity_failures": int(row["integrity_failures"])})
    write_csv(out / "plot_data/fig_results_local_cache_reuse.csv", plot_rows)
    return metadata_base("fig_results_local_cache_reuse", "fig:cache-reuse", "local cache reuse",
                         sources, ["Cache Trials", "Cache Ablation"], [source_path], list(plot_rows[0]), {},
                         "Median origin-side transfer and all independent trial retrieval times with box summaries",
                         len(rows), 0, [row["total_retrieval_time_seconds"] for row in plot_rows], outputs, font)


def accountability_points(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    indexed = {(int(row["fleet_size"]), int(row["batch_size"]), row["reporting_mode"]): row for row in rows}
    result = []
    for size in [100, 500, 1000]:
        for batch in [25, 50, 100, 200]:
            baseline = indexed[(size, batch, "naive_per_device")]
            for mode in ["v1_integrity_root", "v2_complete_witnessed"]:
                row = indexed[(size, batch, mode)]
                result.append({
                    "cohort_size_software_identities": size,
                    "batch_size": batch,
                    "reporting_approach": MODE_LABELS[mode],
                    "projected_transaction_count": int(row["formula_transaction_count"]),
                    "projected_storage_slot_bytes": int(row["formula_storage_slot_bytes"]),
                    "projected_gas": int(row["formula_gas_from_path_median"]),
                    "executed_median_logical_path_gas_basis": int(row["measured_median_gas_basis"]),
                    "storage_relative_to_per_device_percent": 100 * int(row["formula_storage_slot_bytes"]) / int(baseline["formula_storage_slot_bytes"]),
                    "gas_relative_to_per_device_percent": 100 * int(row["formula_gas_from_path_median"]) / int(baseline["formula_gas_from_path_median"]),
                })
    return result


def render_accountability_panel(ax: plt.Axes, rows: list[dict[str, Any]], size: int, metric: str, panel: str) -> None:
    key = "storage_relative_to_per_device_percent" if metric == "storage" else "gas_relative_to_per_device_percent"
    for approach, color, marker, line in [
        ("Integrity-only aggregation", BLUE, "o", "-"),
        ("Complete-cohort witnessed aggregation", ORANGE, "s", "--"),
    ]:
        selected = sorted([row for row in rows if row["cohort_size_software_identities"] == size and row["reporting_approach"] == approach], key=lambda row: row["batch_size"])
        ax.plot([row["batch_size"] for row in selected], [row[key] for row in selected],
                color=color, marker=marker, linestyle=line, markerfacecolor=WHITE, markeredgewidth=1.0, label=approach)
    ax.set_xticks([25, 50, 100, 200])
    ax.set_ylim(0, 11)
    ax.set_xlabel("Batch size")
    ax.set_ylabel("Relative to per-device reporting (%)")
    ax.set_title(f"{size:,} identities: {'storage footprint' if metric == 'storage' else 'aggregate gas'}")
    panel_label(ax, panel)
    style_axis(ax, "y")


def figure_r6(sources: Sources, out: Path, font: str) -> dict[str, Any]:
    source_path, rows = sources.csv_rows("Acct Ablation")
    stats_path = sources.csv_path("Acct Path Statistics")
    totals_path = sources.csv_path("Acct Path Totals")
    break_even_path = sources.csv_path("Gas Break Even")
    stats = {str(row["reporting_mode"]): row for row in sources.sheet("Acct Path Statistics")}
    expected_gas = {"naive_per_device": 169245, "v1_integrity_root": 97836, "v2_complete_witnessed": 434143}
    for mode, expected in expected_gas.items():
        if int(stats[mode]["gas_median"]) != expected:
            raise ValueError(f"Executed median gas mismatch for {mode}")
    points = accountability_points(rows)
    indexed_raw = {(int(row["fleet_size"]), int(row["batch_size"]), row["reporting_mode"]): row for row in rows}
    expected_tx = {"naive_per_device": 500, "v1_integrity_root": 10, "v2_complete_witnessed": 30}
    for mode, count in expected_tx.items():
        if int(indexed_raw[(500, 50, mode)]["formula_transaction_count"]) != count:
            raise ValueError(f"N=500, B=50 transaction mismatch for {mode}")
    outputs: list[Path] = []
    specs = [(size, metric) for size in [100, 500, 1000] for metric in ["storage", "gas"]]
    for index, (size, metric) in enumerate(specs):
        fig, ax = plt.subplots(figsize=(3.45, 2.65))
        render_accountability_panel(ax, points, size, metric, f"({chr(97 + index)})")
        if index == 0:
            ax.legend(frameon=False, fontsize=7, loc="upper right")
        outputs += save_figure(fig, out / f"panels/fig_results_storage_bounded_accountability_{chr(97 + index)}")
    fig, axes = plt.subplots(3, 2, figsize=(7.2, 7.4), sharex=True, sharey=True, layout="constrained")
    for index, (ax, (size, metric)) in enumerate(zip(axes.flat, specs, strict=True)):
        render_accountability_panel(ax, points, size, metric, f"({chr(97 + index)})")
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="outside lower center", ncol=2, frameon=False)
    outputs += save_figure(fig, out / "fig_results_storage_bounded_accountability")
    write_csv(out / "plot_data/fig_results_storage_bounded_accountability.csv", points)
    plotted = [row["storage_relative_to_per_device_percent"] for row in points] + [row["gas_relative_to_per_device_percent"] for row in points]
    return metadata_base("fig_results_storage_bounded_accountability", "fig:accountability-tradeoff",
                         "compiler-layout and executed-path-grounded analytical projections", sources,
                         ["Acct Ablation", "Acct Path Statistics", "Acct Path Totals", "Gas Break Even"],
                         [source_path, stats_path, totals_path, break_even_path], list(points[0]),
                         {"cohort_sizes": [100, 500, 1000], "batch_sizes": [25, 50, 100, 200],
                          "plotted_approaches": ["Integrity-only aggregation", "Complete-cohort witnessed aggregation"],
                          "normalization_baseline": "Per-device reporting = 100%"},
                         "Analytical projection normalized to per-device reporting using compiler-layout bytes and executed median logical-path gas",
                         len(rows), 0, plotted, outputs, font)


def render_validator_liveness(ax: plt.Axes, counts: list[int], panel: str) -> None:
    x = np.arange(4)
    ax.bar(x, counts, width=0.56, color=[BLUE, GREEN, ORANGE, PURPLE], edgecolor=BLACK, linewidth=0.7)
    ax.set_xticks(x, ["4 active", "3 active", "2 active", "Restored\nto 3"])
    ax.set_ylim(0, 5.5)
    ax.set_yticks(range(0, 6))
    ax.set_ylabel("Trials with observed block progress")
    ax.set_title("Crash-fault liveness boundary")
    for position, count in zip(x, counts, strict=True):
        ax.text(position, count + 0.12, f"{count}/5", ha="center", va="bottom", fontsize=8)
    panel_label(ax, panel)
    style_axis(ax, "y")


def render_validator_restoration(ax: plt.Axes, restored: list[dict[str, str]], panel: str) -> None:
    trials = [int(row["trial"]) for row in restored]
    first_block = [float(row["restart_to_first_block_seconds"]) for row in restored]
    first_receipt = [float(row["restart_to_first_receipt_seconds"]) for row in restored]
    ax.plot(trials, first_block, color=BLUE, marker="o", markerfacecolor=WHITE, label="First new block")
    ax.plot(trials, first_receipt, color=ORANGE, marker="s", markerfacecolor=WHITE, linestyle="--", label="First successful receipt")
    ax.set_xticks(trials)
    ax.set_xlabel("Independent trial")
    ax.set_ylabel("Time from validator restart (s)")
    ax.set_title("Quorum restoration")
    ax.legend(frameon=False, loc="best")
    panel_label(ax, panel)
    style_axis(ax, "y")


def figure_r7(sources: Sources, out: Path, font: str) -> dict[str, Any]:
    source_path, rows = sources.csv_rows("Validator Faults")
    cases = ["zero_validators_stopped", "one_validator_stopped", "two_validators_stopped", "quorum_restored_after_validator_restart"]
    groups = {case: [row for row in rows if row["fault_case"] == case] for case in cases}
    if len(rows) != 20 or any(len(group) != 5 for group in groups.values()):
        raise ValueError("Validator campaign must contain five trials and four phases per trial")
    counts = [sum(row["observed_progress"] == "True" for row in groups[case]) for case in cases]
    if counts != [5, 5, 0, 5]:
        raise ValueError(f"Unexpected validator progress counts: {counts}")
    restored = sorted(groups[cases[-1]], key=lambda row: int(row["trial"]))
    outputs: list[Path] = []
    fig, ax = plt.subplots(figsize=(3.45, 2.8)); render_validator_liveness(ax, counts, "(a)")
    outputs += save_figure(fig, out / "panels/fig_results_validator_crash_faults_a")
    fig, ax = plt.subplots(figsize=(3.45, 2.8)); render_validator_restoration(ax, restored, "(b)")
    outputs += save_figure(fig, out / "panels/fig_results_validator_crash_faults_b")
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.15), layout="constrained")
    render_validator_liveness(axes[0], counts, "(a)"); render_validator_restoration(axes[1], restored, "(b)")
    outputs += save_figure(fig, out / "fig_results_validator_crash_faults")
    plot_rows = [{"trial": int(row["trial"]), "phase": row["fault_case"],
                  "active_validator_processes": int(row["active_validator_processes"]),
                  "observed_block_progress": row["observed_progress"] == "True",
                  "restart_to_first_block_seconds": float(row["restart_to_first_block_seconds"]) if row["restart_to_first_block_seconds"] else None,
                  "restart_to_first_receipt_seconds": float(row["restart_to_first_receipt_seconds"]) if row["restart_to_first_receipt_seconds"] else None}
                 for row in rows]
    write_csv(out / "plot_data/fig_results_validator_crash_faults.csv", plot_rows)
    plotted = [float(row["restart_to_first_block_seconds"]) for row in restored] + [float(row["restart_to_first_receipt_seconds"]) for row in restored]
    return metadata_base("fig_results_validator_crash_faults", "fig:validator-faults", "local validator crash-fault experiment",
                         sources, ["Validator Faults"], [source_path], list(plot_rows[0]), {},
                         "Observed-progress trial counts and all per-trial quorum-restoration times",
                         len(rows), 0, plotted, outputs, font)


FIGURES: list[tuple[str, Callable[[Sources, Path, str], dict[str, Any]]]] = [
    ("fig_results_block_period_sensitivity", figure_r1),
    ("fig_results_success_share_software_cohorts", figure_r2),
    ("fig_results_failure_probability_sensitivity", figure_r3),
    ("fig_results_application_shaped_ipfs_retrieval", figure_r4),
    ("fig_results_local_cache_reuse", figure_r5),
    ("fig_results_storage_bounded_accountability", figure_r6),
    ("fig_results_validator_crash_faults", figure_r7),
]


CAPTIONS = {
    "fig_results_block_period_sensitivity": ("Submission-to-receipt latency under 1-, 2-, and 4-second configured block periods. Each configuration contains 30 transactions. Individual observations and box summaries are shown on a logarithmic time axis; the retained approximately 31.032-second observation in the 1-second profile remains visible.", "fig:block-period-sensitivity", "Block-Period Sensitivity"),
    "fig_results_success_share_software_cohorts": ("Mean success share across authenticated software-cohort configurations. Each point summarizes 20 deterministic seeds for the stated cohort size and configured failure probability; error bars denote bootstrap 95\\% confidence intervals.", "fig:fleet-success-share", "Main software-fleet outcome table"),
    "fig_results_failure_probability_sensitivity": ("Outcome sensitivity for a 500-identity software cohort. Panels show the mean shares of (a) SUCCESS, (b) ROLLBACK, and (c) FAIL outcomes across 20 deterministic seeds; error bars denote bootstrap 95\\% confidence intervals.", "fig:failure-probability-sensitivity", "500-identity outcome-sensitivity paragraph"),
    "fig_results_application_shaped_ipfs_retrieval": ("Application-shaped local IPFS retrieval time for a 537,088-byte artifact. Each profile contains 20 completed transfers shown on a logarithmic time axis. Application-level delays and paced reads were used; all observations were retained, including the 915.737-second observation in P3.", "fig:ipfs-retrieval", "Retrieval-distribution table"),
    "fig_results_local_cache_reuse": ("Local cache-reuse ablation across ten independent trials per mode, with 1,000 sequential requests in each trial. Panel (a) reports origin-side transfer, whereas panel (b) shows the distribution of total retrieval time. Each cache-enabled trial began with an empty in-process payload buffer and performed one origin retrieval before serving subsequent requests from the buffer. No integrity failures occurred.", "fig:cache-reuse", "Cache summary table"),
    "fig_results_storage_bounded_accountability": ("Batch-size sensitivity of integrity-only and complete-cohort witnessed aggregation. The left column reports projected on-chain storage footprint and the right column reports projected aggregate gas, each normalized to per-device reporting at 100\\%. Rows correspond to cohorts of 100, 500, and 1,000 software identities. Storage values use compiler-reported layouts, whereas gas projections use the executed median logical-path gas for each reporting approach. The per-device 100\\% baseline is omitted from the panels to preserve resolution among the aggregation curves.", "fig:accountability-tradeoff", "Analytical accountability-sensitivity table"),
    "fig_results_validator_crash_faults": ("Local validator crash-fault and quorum-restoration results across five independent trials. Panel (a) reports observed block progress with four, three, and two active validators and after restoring a third validator. Panel (b) reports the time from validator restart to the first new block and first successful transaction receipt. The campaign injected process crashes only.", "fig:validator-faults", "Crash-fault and restoration table"),
}


def latex_figure(name: str) -> str:
    caption, label, placement = CAPTIONS[name]
    multi = {
        "fig_results_failure_probability_sensitivity": 3,
        "fig_results_local_cache_reuse": 2,
        "fig_results_storage_bounded_accountability": 6,
        "fig_results_validator_crash_faults": 2,
    }
    lines = [f"% Recommended placement: after {placement}.", "\\begin{figure*}[t]" if name in multi else "\\begin{figure}[t]", "  \\centering"]
    if name not in multi:
        lines.append(f"  \\includegraphics[width=\\linewidth]{{paper_assets/figures/results/{name}.pdf}}")
    else:
        count = multi[name]
        width = "0.32\\textwidth" if count == 3 else "0.48\\textwidth"
        for index in range(count):
            suffix = chr(97 + index)
            lines.extend([f"  \\begin{{minipage}}[t]{{{width}}}", "    \\centering",
                          f"    \\includegraphics[width=\\linewidth]{{paper_assets/figures/results/panels/{name}_{suffix}.pdf}}",
                          "  \\end{minipage}%" if (index + 1) % (3 if count == 3 else 2) else "  \\end{minipage}\\par\\smallskip"])
    lines.extend([f"  \\caption{{{caption}}}", f"  \\label{{{label}}}", "\\end{figure*}" if name in multi else "\\end{figure}", ""])
    return "\n".join(lines)


def write_latex(out: Path) -> None:
    content = "% Generated Results figure environments. Requires the manuscript's existing graphicx support.\n\n"
    content += "\n".join(latex_figure(name) for name, _ in FIGURES)
    (out / "results_figures.tex").write_text(content, encoding="utf-8")


def write_inventory(out: Path, metadata: list[dict[str, Any]]) -> None:
    purposes = {
        "fig_results_block_period_sensitivity": "Block-period sensitivity",
        "fig_results_success_share_software_cohorts": "Success share across software cohorts",
        "fig_results_failure_probability_sensitivity": "Outcome sensitivity at 500 identities",
        "fig_results_application_shaped_ipfs_retrieval": "Application-shaped local IPFS retrieval distribution",
        "fig_results_local_cache_reuse": "Origin transfer and retrieval-time cache ablation",
        "fig_results_storage_bounded_accountability": "Projected storage and gas trade-off",
        "fig_results_validator_crash_faults": "Crash-fault liveness boundary and restoration",
    }
    lines = ["# Results Figure Inventory", "", "| Figure | Placement | Scientific purpose | Evidence class | Source | Sample size | Principal statistic | PDF | SVG | PNG | LaTeX label | Status |", "|---|---|---|---|---|---:|---|---|---|---|---|---|"]
    for item in metadata:
        name = item["figure_semantic_name"]
        _, _, placement = CAPTIONS[name]
        paths = {Path(record["path"]).suffix: record["path"] for record in item["output_files"] if "/panels/" not in record["path"]}
        lines.append(f"| `{name}` | After {placement} | {purposes[name]} | {item['evidence_class']} | {', '.join(item['source_worksheet_names'])} | {item['row_count']} | {item['aggregation_method']} | `{paths['.pdf']}` | `{paths['.svg']}` | `{paths['.png']}` | `{item['manuscript_label']}` | Pending validation |")
    (out / "figure_inventory.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def create_contact_sheet(out: Path) -> None:
    files = [out / f"{name}.png" for name, _ in FIGURES]
    thumbs = []
    for file in files:
        image = Image.open(file).convert("RGB")
        image.thumbnail((1400, 900), Image.Resampling.LANCZOS)
        thumbs.append((file.stem, image.copy()))
    width, cell_height = 1500, 1020
    canvas = Image.new("RGB", (width * 2, cell_height * 4), "white")
    draw = ImageDraw.Draw(canvas)
    for index, (name, image) in enumerate(thumbs):
        col, row = index % 2, index // 2
        x = col * width + (width - image.width) // 2
        y = row * cell_height + 65
        canvas.paste(image, (x, y))
        draw.text((col * width + 30, row * cell_height + 20), name, fill="black")
    path = out / "previews/results_figure_contact_sheet.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(path, dpi=(300, 300))


def clean_generated(out: Path) -> None:
    if out.exists():
        shutil.rmtree(out)
    print(f"Removed generated figure root: {out}")


def validate_only(workbook: Path, out: Path) -> None:
    command = [sys.executable, str(ROOT / "experiments/analysis/validate_results_figures.py"),
               "--workbook", str(workbook), "--figure-root", str(out)]
    raise SystemExit(subprocess.call(command, cwd=ROOT))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=ROOT / "LedgerGuard_Comprehensive_Results.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "paper_assets/figures/results")
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--clean-generated", action="store_true")
    args = parser.parse_args()
    workbook_path = args.workbook if args.workbook.is_absolute() else ROOT / args.workbook
    output = args.output if args.output.is_absolute() else ROOT / args.output
    if args.clean_generated:
        clean_generated(output)
        return
    if args.validate_only:
        validate_only(workbook_path, output)
        return
    font = select_font()
    configure_style(font)
    output.mkdir(parents=True, exist_ok=True)
    for child in ["panels", "previews", "plot_data", "metadata"]:
        (output / child).mkdir(parents=True, exist_ok=True)
    sources = Sources(workbook_path)
    metadata = []
    for name, builder in FIGURES:
        item = builder(sources, output, font)
        (output / f"metadata/{name}.json").write_text(json.dumps(item, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        metadata.append(item)
        print(f"generated {name}")
    write_latex(output)
    write_inventory(output, metadata)
    create_contact_sheet(output)
    print(f"Generated {len(metadata)} main Results figures using {font}.")


if __name__ == "__main__":
    main()
