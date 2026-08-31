#!/usr/bin/env python3
"""Create a workbook containing the strong-evaluation CSV outputs."""

from __future__ import annotations

import argparse
import csv
import json
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

STRONG_EVAL_CSVS = [
    "sqlite_baseline_timing.csv",
    "governance_timing.csv",
    "governance_correctness.csv",
    "distributed_besu_timing.csv",
    "validator_resource_usage.csv",
    "chain_growth.csv",
    "network_sensitivity.csv",
    "edge_cache_ablation.csv",
    "accountability_ablation.csv",
    "adversarial_validation.csv",
    "security_checks.csv",
]

def csv_rows(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-dir", default="results")
    args = parser.parse_args()

    results_dir = Path(args.results_dir)
    csv_dir = results_dir / "csv"
    workbook_dir = results_dir / "workbook"
    workbook_dir.mkdir(parents=True, exist_ok=True)
    out_path = workbook_dir / "ledgerguard_strong_eval_results.xlsx"

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        payload = {name[:-4]: csv_rows(csv_dir / name) for name in STRONG_EVAL_CSVS if (csv_dir / name).exists()}
        _write_minimal_xlsx(out_path, payload)
        sidecar = workbook_dir / "ledgerguard_strong_eval_results.sidecar.json"
        sidecar.write_text(json.dumps({"note": "Workbook written with standard-library XLSX fallback because openpyxl is unavailable."}, indent=2) + "\n", encoding="utf-8")
        print(out_path)
        return 0

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    for name in STRONG_EVAL_CSVS:
        csv_path = csv_dir / name
        if not csv_path.exists():
            continue
        ws = wb.create_sheet(csv_path.stem[:31])
        for row in csv_rows(csv_path):
            ws.append(row)
        if ws.max_row:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = min(max(len(str(cell.value or "")) for cell in col) + 2, 42)
                ws.column_dimensions[col[0].column_letter].width = width

    if not wb.sheetnames:
        ws = wb.create_sheet("README")
        ws.append(["No CSV results were available."])

    wb.save(out_path)
    print(out_path)
    return 0


def _sheet_xml(rows: list[list[str]]) -> str:
    body = []
    for r_idx, row in enumerate(rows, start=1):
        cells = []
        for c_idx, value in enumerate(row, start=1):
            ref = f"{_col(c_idx)}{r_idx}"
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(str(value))}</t></is></c>')
        body.append(f'<row r="{r_idx}">{"".join(cells)}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{"".join(body)}</sheetData></worksheet>'
    )


def _col(index: int) -> str:
    result = ""
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def _write_minimal_xlsx(path: Path, sheets: dict[str, list[list[str]]]) -> None:
    if not sheets:
        sheets = {"README": [["No CSV results were available."]]}
    sheet_names = [name[:31] or "Sheet" for name in sheets]
    workbook_sheets = "".join(
        f'<sheet name="{escape(name)}" sheetId="{idx}" r:id="rId{idx}"/>'
        for idx, name in enumerate(sheet_names, start=1)
    )
    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<sheets>{workbook_sheets}</sheets></workbook>'
    )
    rels = "".join(
        f'<Relationship Id="rId{idx}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{idx}.xml"/>'
        for idx in range(1, len(sheet_names) + 1)
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/></Types>')
        zf.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        zf.writestr("xl/workbook.xml", workbook_xml)
        zf.writestr("xl/_rels/workbook.xml.rels", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>')
        for idx, (_, rows) in enumerate(sheets.items(), start=1):
            zf.writestr(f"xl/worksheets/sheet{idx}.xml", _sheet_xml(rows))


if __name__ == "__main__":
    raise SystemExit(main())
